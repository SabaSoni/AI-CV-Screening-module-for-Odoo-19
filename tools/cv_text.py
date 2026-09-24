# -*- coding: utf-8 -*-
"""Plain-text extraction from CV files. Pure Python, no Odoo imports.

Supported: PDF (pypdf), DOCX (python-docx), TXT/MD. Legacy binary .doc is
detected and reported; the caller can fall back to Odoo's own index_content.
"""
import io
import re
import unicodedata

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency
    try:
        from PyPDF2 import PdfReader
    except Exception:
        PdfReader = None

try:
    import docx  # python-docx
except Exception:  # pragma: no cover - optional dependency
    docx = None

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def file_kind(name, mimetype=None):
    """Classify a file as 'pdf', 'docx', 'txt', 'doc' or None."""
    n = (name or "").lower().strip()
    m = (mimetype or "").lower().strip()
    if n.endswith(".pdf") or m == "application/pdf":
        return "pdf"
    if n.endswith(".docx") or m == DOCX_MIME:
        return "docx"
    if n.endswith((".txt", ".md")) or m in ("text/plain", "text/markdown"):
        return "txt"
    if n.endswith(".doc") or m == "application/msword":
        return "doc"
    return None


def looks_like_cv_file(name, mimetype=None):
    return file_kind(name, mimetype) is not None


def clean_text(text):
    """Normalise unicode, line endings and whitespace; keep single blank lines."""
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.split("\n")]
    out, blank = [], 0
    for line in lines:
        if line:
            out.append(line)
            blank = 0
        else:
            blank += 1
            if blank <= 1:
                out.append("")
    return "\n".join(out).strip()


def extract_text(name, data, mimetype=None):
    """Return (text, note). `note` is a human-readable remark, e.g. OCR needed."""
    kind = file_kind(name, mimetype)
    if kind == "pdf":
        text = _extract_pdf(data)
    elif kind == "docx":
        text = _extract_docx(data)
    elif kind == "txt":
        text = _extract_txt(data)
    elif kind == "doc":
        return "", "Legacy .doc format is not supported; re-save as .docx or PDF."
    else:
        return "", "Unsupported file type (use PDF, DOCX or TXT)."
    text = clean_text(text)
    note = ""
    if kind == "pdf" and len(text) < 80:
        note = "PDF has almost no extractable text; it is probably a scanned image (OCR needed)."
    return text, note


def _extract_pdf(data):
    if PdfReader is None:
        raise RuntimeError("pypdf is not installed (pip install pypdf)")
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            # One unreadable page must not kill the whole CV.
            parts.append("")
    return "\n".join(parts)


def _extract_docx(data):
    if docx is None:
        raise RuntimeError("python-docx is not installed (pip install python-docx)")
    document = docx.Document(io.BytesIO(data))
    parts = []
    # Headers often hold the contact block, so read them first.
    for section in document.sections:
        try:
            parts.extend(p.text for p in section.header.paragraphs if p.text.strip())
        except Exception:
            pass
    parts.extend(p.text for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:  # merged cells repeat their text
                    cells.append(t)
            if cells:
                parts.append(" | ".join(cells))
    for section in document.sections:
        try:
            parts.extend(p.text for p in section.footer.paragraphs if p.text.strip())
        except Exception:
            pass
    return "\n".join(parts)


def _extract_txt(data):
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")

# -*- coding: utf-8 -*-
"""Stand-alone checks for the engine (no Odoo, no database).

Run with the Python that runs Odoo, so the same libraries are used:
    Windows:  "C:\\Program Files\\Odoo 19.0.<build>\\python\\python.exe" tools\\selftest.py
    Linux:    python3 tools/selftest.py

Optional: pass a CV file path to see what the engine extracts from it.
If Ollama is running, a live LLM + embedding smoke test is executed too.
"""
import os
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))  # addon root, so `tools.*` imports work

from tools import contacts, cv_text, matcher, prompts  # noqa: E402
from tools.ollama_client import OllamaClient  # noqa: E402

SAMPLE_CV = """
Nino Beridze
Senior Backend Developer
Tbilisi, Georgia | Tel: +995 555 12 34 56 | nino.beridze@example.com
linkedin.com/in/nino-beridze

Experience
2019 - 2024  Backend Developer, FMG Soft
  Built REST APIs with Python (Django, FastAPI) and PostgreSQL. Deployed with Docker on Linux.
  Led a team of 3 developers. Odoo module development for HR and accounting.

Education
2015 - 2019  BSc Computer Science, Tbilisi State University

Languages: Georgian (native), English (C1), Russian (B2)
"""

REQUIREMENTS = [
    {"id": 1, "name": "Python backend", "keywords": ["python", "django", "fastapi", "flask"], "weight": 5, "mandatory": True},
    {"id": 2, "name": "PostgreSQL", "keywords": ["postgresql", "postgres"], "weight": 3, "mandatory": False},
    {"id": 3, "name": "Odoo development", "keywords": ["odoo"], "weight": 4, "mandatory": False},
    {"id": 4, "name": "Kubernetes", "keywords": ["kubernetes", "k8s"], "weight": 2, "mandatory": False},
    {"id": 5, "name": "English B2+", "keywords": ["english"], "weight": 1, "mandatory": False},
]

failures = 0


def check(label, condition, detail=""):
    global failures
    status = "OK  " if condition else "FAIL"
    if not condition:
        failures += 1
    print("[%s] %s%s" % (status, label, (" -> " + str(detail)) if detail else ""))


def test_contacts():
    c = contacts.extract_contacts(SAMPLE_CV)
    check("email", c["email"] == "nino.beridze@example.com", c["email"])
    check("phone (E.164)", c["phone"] == "+995555123456", c["phone"])
    check("linkedin", c["linkedin"] == "https://linkedin.com/in/nino-beridze", c["linkedin"])
    check("name", c["name"] == "Nino Beridze", c["name"])
    # Years must not be mistaken for phone numbers.
    c2 = contacts.extract_contacts("Experience 2015 - 2019 and 2019-2024\nMob: 599 11 22 33")
    check("phone ignores year ranges", c2["phone"] == "+995599112233", c2["phone"])
    ka = contacts.extract_contacts("ნინო ბერიძე\nპროგრამისტი\nტელ: 577 00 11 22\nnino@example.ge")
    check("georgian name", ka["name"] == "ნინო ბერიძე", ka["name"])
    check("georgian phone", ka["phone"] == "+995577001122", ka["phone"])
    labeled = contacts.extract_contacts("CURRICULUM VITAE\nსახელი და გვარი: გიორგი მაისურაძე\nEmail: g@example.ge")
    check("labeled georgian name", labeled["name"] == "გიორგი მაისურაძე", labeled["name"])
    check("job title is not a name", contacts.guess_name("Senior Developer\nSkills: Python") == "")
    multi = contacts.extract_contacts("Mob: 577 44 55 66" + chr(10) + "2016 - 2024 Chief accountant")
    check("phone does not swallow the next line", multi["phone"] == "+995577445566", multi["phone"])


def test_matcher():
    norm = matcher.normalize(SAMPLE_CV)
    layer = {}
    for r in REQUIREMENTS:
        hits = matcher.keyword_hits(norm, r["keywords"])
        if hits:
            layer[r["id"]] = hits
    check("keyword hits", set(layer) == {1, 2, 3, 5}, sorted(layer))
    check("'c' does not match inside words", not matcher.keyword_hits(matcher.normalize("I like cats"), ["c"]))
    check("'c++' matches", bool(matcher.keyword_hits(matcher.normalize("Skills: C++, Java"), ["c++"])))
    check("georgian suffix tolerated", bool(matcher.keyword_hits(matcher.normalize("პითონის ცოდნა"), ["პითონი"])))

    combined = matcher.combine(REQUIREMENTS, layer)
    # weights: 5+3+4+1 = 13 of 15 -> 87 %
    check("weighted percent", combined["percent"] == 87, combined["percent"])
    check("mandatory ok", combined["mandatory_missing"] is False)
    check("verdict match @80", matcher.verdict(87, 80, False) == "match")
    check("verdict partial @90", matcher.verdict(87, 90, False) == "partial")
    check("verdict partial when must-have missing", matcher.verdict(95, 80, True) == "partial")
    check("verdict no_match", matcher.verdict(40, 80, False) == "no_match")

    llm = {4: {"status": "partial", "evidence": "Docker on Linux, no k8s"}}
    combined2 = matcher.combine(REQUIREMENTS, layer, None, llm)
    k8s = [r for r in combined2["results"] if r["id"] == 4][0]
    check("llm partial applied", k8s["status"] == "partial" and "llm" in k8s["sources"], k8s)
    chunks = matcher.chunk_text("word " * 1000, size=700, overlap=150)
    check("chunking", len(chunks) > 5 and all(len(c) <= 700 for c in chunks), len(chunks))


def test_evidence_policy():
    """The model's opinion only counts when the CV backs it (a real false 'match' happened:
    a CV without the word Odoo got credit for 'Odoo development')."""
    from tools import experience

    cv = "Giorgi Kapanadze\n2021 - 2024 Backend developer: Python, Flask, PostgreSQL, Docker.\nLed a team of 3 developers."
    norm = matcher.normalize(cv)
    reqs = [
        {"id": 1, "name": "Odoo development", "keywords": ["odoo"], "explicit_keywords": True, "weight": 4},
        {"id": 2, "name": "Team leadership", "keywords": ["Team leadership"], "explicit_keywords": False, "weight": 2},
        {"id": 3, "name": "Python", "keywords": ["python"], "explicit_keywords": True, "weight": 5},
        {"id": 4, "name": "Docker", "keywords": ["containers"], "explicit_keywords": True, "weight": 1},
    ]
    keyword_layer = {r["id"]: matcher.keyword_hits(norm, r["keywords"]) for r in reqs}
    llm = {
        1: {"status": "partial", "quote": "Odoo module development for HR", "evidence": "invented"},
        2: {"status": "met", "quote": "Led a team of 3 developers", "evidence": "led a team"},
        3: {"status": "partial", "quote": "", "evidence": "python"},
        4: {"status": "met", "quote": "Backend developer: Python, Flask, PostgreSQL, Docker", "evidence": "docker"},
    }
    fixed = matcher.apply_evidence_policy(reqs, keyword_layer, {}, llm, norm)
    check("invented claim (quote not in the CV) does not count",
          fixed[1]["status"] == "not_met" and fixed[1]["policy"] == "unverified", fixed[1])
    check("verified quote on a requirement without keywords counts fully", fixed[2]["status"] == "met", fixed[2])
    check("keyword found: the model cannot lower it", fixed[3]["status"] == "partial" and not fixed[3].get("policy"))
    check("keywords typed but missing: verified 'met' is worth half",
          fixed[4]["status"] == "partial" and fixed[4]["policy"] == "keywords_missing", fixed[4])
    combined = matcher.combine(reqs, keyword_layer, {}, fixed)
    # python 5 (keyword) + leadership 2 + docker 0.5 = 7.5 of 12 -> 62 %
    check("score after the policy", combined["percent"] == 62, combined["percent"])
    check("quote check tolerates punctuation and case", matcher.quote_in_text('"led a TEAM of 3 developers."', norm))
    check("quote check rejects text that is not there", not matcher.quote_in_text("five years of Odoo", norm))

    years, spans = experience.estimate_years(
        "Experience\n2019 - 2021 Developer, A\n03/2020 - 08/2023 Developer, B\nJan 2024 - present Lead, C\n"
        "Education\n2012 - 2016 Tbilisi State University\n2016 - 2018 GTU", 2026)
    check("experience from dates: overlaps merged, education skipped", years == 6 and spans == [(2019, 2023), (2024, 2026)],
          "%s %s" % (years, spans))
    check("education line is skipped without a heading",
          experience.estimate_years("2015 - 2019 BSc Computer Science, University", 2026)[0] == 0)


def test_cv_text():
    text, note = cv_text.extract_text("cv.txt", SAMPLE_CV.encode("utf-8"))
    check("txt extraction", "Nino Beridze" in text and not note)
    try:
        import docx
        d = docx.Document()
        d.sections[0].header.paragraphs[0].text = "Header: nino.beridze@example.com"
        d.add_paragraph("Nino Beridze")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Skill"
        t.rows[0].cells[1].text = "Python"
        import io
        buf = io.BytesIO()
        d.save(buf)
        text, note = cv_text.extract_text("cv.docx", buf.getvalue())
        check("docx extraction (header + table)", "nino.beridze@example.com" in text and "Skill | Python" in text, text[:80])
    except ImportError:
        check("python-docx installed", False)
    # Optional real-PDF check: set ODOO_ADDONS to Odoo's addons folder to use its demo CV.
    demo_pdf = os.path.join(os.environ.get("ODOO_ADDONS", ""), "hr_recruitment", "data",
                            "hr_recruitment_demo_jones_cv.pdf")
    if os.path.exists(demo_pdf):
        with open(demo_pdf, "rb") as fh:
            text, note = cv_text.extract_text("jones.pdf", fh.read())
        c = contacts.extract_contacts(text)
        check("odoo demo pdf extracts", len(text) > 100 and c["email"] == "EnriqueJones@info.com" and c["name"] == "Enrique Jones",
              "%d chars, email=%s, phone=%s, name=%s" % (len(text), c["email"], c["phone"], c["name"]))
    check("unsupported .doc reported", cv_text.extract_text("old.doc", b"x")[1] != "")


def test_prompts():
    schema = prompts.response_schema()
    check("schema has requirements", "requirements" in schema["properties"])
    user = prompts.build_user_prompt("Backend Dev", "Build APIs", "", REQUIREMENTS, "must start in October", SAMPLE_CV)
    check("prompt contains ids and notes", "id 1:" in user and "October" in user)
    subj, body = prompts.fallback_email("ka", "ნინო", "Backend Developer", "FMG Soft")
    check("georgian fallback email", "ნინო" in body and "FMG Soft" in body, subj)


def test_textfix():
    from tools import textfix

    def geo(*codes):  # Georgian words from code points: no typing slips in this file
        return "".join(chr(c) for c in codes)

    cv = "Nino Beridze\nJos" + chr(0x00E9) + " works at FMG Soft"
    nino = geo(0x10DC, 0x10D8, 0x10DC, 0x10DD)
    mangled = geo(0x10D1, 0x10D4, 0x10E0, 0x10EB, 0x10D8, 0x10EB, 0x10D4)      # "berdzidze"
    aris = geo(0x10D0, 0x10E0, 0x10D8, 0x10E1)
    tamil = geo(0x10D1, 0x10D4) + chr(0x0B8E) + chr(0x0BA9)
    cyrillic = chr(0x043F) + chr(0x043B) + chr(0x044E) + chr(0x0441)
    text = "%s %s, %s Servi%sos %s" % (aris, tamil, cyrillic, chr(0x00E7), aris)
    cleaned = textfix.strip_foreign_script(text, "ka", cv)
    check("glitch words removed (Tamil, Cyrillic, Portuguese)", cleaned == "%s, %s" % (aris, aris), cleaned)
    kept = "Jos" + chr(0x00E9) + " " + aris
    check("accented name from the CV is kept", textfix.strip_foreign_script(kept, "ka", cv) == kept)
    fixed = textfix.restore_names("%s%s %s %s" % (nino, geo(0x10E1), mangled, aris), "Nino Beridze")
    check("mangled candidate name restored", fixed == "Nino-%s Beridze %s" % (geo(0x10E1), aris), fixed)

    # The same gap said three ways (a real case reported by the user) must collapse to one.
    tail = "Kubernetes, " + aris + " " + nino
    gaps = ["Nino " + tail + " one", "His profile: " + tail + " two", "Nino again: " + tail + " three",
            "No Docker " + aris]
    merged = textfix.dedupe_similar(gaps, same_topic=True, ignore=["Nino", "Beridze"])
    check("repeated gaps about one technology are merged", len(merged) == 2, merged)
    check("mixed-script glitch word detected",
          textfix.foreign_words(aris + " " + geo(0x10DB, 0x10DD) + "fatter", "ka", "") != [])
    check("hyphenated Latin stem + Georgian suffix is fine",
          textfix.foreign_words("Python-" + geo(0x10D8, 0x10D7) + " Odoo-" + geo(0x10E1), "ka", "") == [])
    check("english text detected as wrong language", textfix.script_ratio("She is a developer", "ka") < 0.35)


def test_ollama_live(chat_model="gemma3:4b", embed_model="bge-m3"):
    client = OllamaClient()
    if not client.is_available():
        print("[SKIP] Ollama not running - live test skipped")
        return
    names = client.list_models()
    print("       Ollama models:", ", ".join(names) or "(none)")
    if client.has_model(embed_model, names):
        t0 = time.time()
        vecs = client.embed(embed_model, ["Python backend development", "Experience with Django and FastAPI", "Fluent in French"])
        sim_related = matcher.cosine(vecs[0], vecs[1])
        sim_unrelated = matcher.cosine(vecs[0], vecs[2])
        check("embeddings: related > unrelated", sim_related > sim_unrelated,
              "related=%.2f unrelated=%.2f (%.1fs)" % (sim_related, sim_unrelated, time.time() - t0))
    else:
        print("[SKIP] embedding model %s not pulled" % embed_model)
    if client.has_model(chat_model, names):
        t0 = time.time()
        out, raw = client.chat_json(
            chat_model,
            prompts.system_prompt("ka", "FMG Soft"),
            prompts.build_user_prompt("Backend Developer", "Python APIs for HR software", "", REQUIREMENTS, "", SAMPLE_CV, language_code="ka"),
            schema=prompts.response_schema(), num_ctx=8192,
        )
        secs = time.time() - t0
        statuses = {int(r["id"]): r["status"] for r in out.get("requirements", []) if "id" in r}
        check("llm: json with candidate + requirements", "candidate" in out and len(statuses) >= 4,
              "%.1fs, score=%s, statuses=%s" % (secs, out.get("overall_score"), statuses))
        check("llm: python met", statuses.get(1) == "met")
        check("llm: kubernetes not met/partial", statuses.get(4) in ("not_met", "partial"))
        print("       summary:", (out.get("summary") or "")[:300].replace("\n", " "))
        print("       email_subject:", out.get("email_subject"))
    else:
        print("[SKIP] chat model %s not pulled" % chat_model)


def show_file(path):
    with open(path, "rb") as fh:
        data = fh.read()
    text, note = cv_text.extract_text(os.path.basename(path), data)
    print("=== %s: %d chars %s" % (path, len(text), ("(" + note + ")") if note else ""))
    print(text[:1500])
    print("=== contacts:", contacts.extract_contacts(text))


if __name__ == "__main__":
    if len(sys.argv) > 1:
        show_file(sys.argv[1])
    test_contacts()
    test_matcher()
    test_cv_text()
    test_prompts()
    test_textfix()
    test_evidence_policy()
    test_ollama_live()
    print("\n%s" % ("ALL CHECKS PASSED" if not failures else "%d CHECK(S) FAILED" % failures))
    sys.exit(1 if failures else 0)
